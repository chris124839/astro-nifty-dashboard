import os
import re
import time
import openpyxl

from bs4 import BeautifulSoup
from datetime import datetime
from dateutil import parser

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC


# =========================================================
# SETTINGS
# =========================================================

GEONAME_ID = "1254360"

FILE = "data/Planets Transists.xlsx"

PLANETS = {
    "Mercury": "budha-transit-date-time",
    "Venus": "shukra-transit-date-time",
    "Mars": "mangal-transit-date-time",
    "Jupiter": "guru-transit-date-time",
    "Saturn": "shani-transit-date-time",
    "Sun": "surya-transit-date-time",
    "Moon": "chandra-transit-date-time",
}


# =========================================================
# FIND LAST YEAR IN EXISTING FILE
# =========================================================

if not os.path.exists(FILE):
    raise FileNotFoundError(
        f"{FILE} not found. Upload the existing Planet Transits file first."
    )

wb = openpyxl.load_workbook(FILE)

if "02_Raw_Transits" not in wb.sheetnames:
    raise ValueError(
        "Sheet '02_Raw_Transits' not found."
    )

ws = wb["02_Raw_Transits"]

last_date = None

for row in ws.iter_rows(min_row=2, values_only=True):

    value = row[0]

    if not value:
        continue

try:

    if isinstance(value, datetime):
        dt = value

    else:
        text = str(value).strip()

        try:
            dt = datetime.fromisoformat(
                text.replace("Z", "")
            )

        except:
            dt = parser.parse(
                text,
                dayfirst=True
            )

        if last_date is None or dt > last_date:
            last_date = dt

    except Exception:
        continue


if last_date is None:
    raise ValueError(
        "Could not find a valid date in the existing Excel file."
    )


# =========================================================
# NEXT FULL YEAR
# =========================================================

next_year = last_date.year + 1

print("Existing last date:", last_date.strftime("%d/%m/%Y"))
print("Next year to scrape:", next_year)


# =========================================================
# CHROME
# =========================================================

options = Options()

options.add_argument("--headless=new")
options.add_argument("--no-sandbox")
options.add_argument("--disable-dev-shm-usage")
options.add_argument("--disable-gpu")
options.add_argument("--window-size=1920,1080")
options.add_argument("--lang=en-US")

driver = webdriver.Chrome(options=options)

wait = WebDriverWait(driver, 20)


# =========================================================
# EXISTING RECORDS
# =========================================================

existing = set()

for row in ws.iter_rows(min_row=2, values_only=True):

    if not row[0]:
        continue

    key = tuple(str(x) if x is not None else "" for x in row[:5])

    existing.add(key)


# =========================================================
# SCRAPE NEXT YEAR
# =========================================================

try:

    for planet, slug in PLANETS.items():

        print()
        print("=" * 60)
        print(f"{planet} - {next_year}")
        print("=" * 60)

        url = (
            f"https://www.drikpanchang.com/planet/transit/"
            f"{slug}.html"
            f"?year={next_year}"
            f"&geoname-id={GEONAME_ID}"
            f"&lang=en"
        )

        print(url)

        driver.get(url)

        wait.until(
            EC.presence_of_element_located(
                (
                    By.CSS_SELECTOR,
                    ".dpCard, .dpHighlightedCard"
                )
            )
        )

        time.sleep(2)

        soup = BeautifulSoup(
            driver.page_source,
            "html.parser"
        )

        cards = soup.select(
            ".dpCard, .dpHighlightedCard"
        )

        added = 0

        for card in cards:

            title = card.find(
                "div",
                class_="dpTitle"
            )

            value = card.find(
                "div",
                class_="dpValue"
            )

            if not title or not value:
                continue

            raasi = title.get_text(
                " ",
                strip=True
            )

            raw = value.get_text(
                " ",
                strip=True
            )

            # ---------------------------------------------
            # Retrograde
            # ---------------------------------------------

            retrograde = (
                "Retrograde"
                if (
                    "\u21ba" in raw
                    or "\u21bb" in raw
                    or "Retrograde" in raw
                )
                else "Forward"
            )

            motion = retrograde

            # ---------------------------------------------
            # Clean
            # ---------------------------------------------

            clean = raw.replace(
                "\u21ba",
                " "
            ).replace(
                "\u21bb",
                " "
            )

            clean = re.sub(
                r"\s+",
                " ",
                clean
            ).strip()

            # ---------------------------------------------
            # Extract date/time
            # ---------------------------------------------

            pattern = (
                r"([A-Za-z]+\s+\d{1,2},\s+\d{4}"
                r".*?"
                r"\d{1,2}:\d{2})"
            )

            match = re.search(
                pattern,
                clean
            )

            if not match:
                continue

            try:

                dt = parser.parse(
                    match.group(1)
                )

            except Exception:
                continue

            # Only accept requested year

            if dt.year != next_year:
                continue

            date_value = dt.strftime(
                "%d/%m/%Y"
            )

            time_value = dt.strftime(
                "%H:%M"
            )

            row = (
                date_value,
                planet,
                raasi,
                time_value,
                retrograde,
                motion
            )

            # ---------------------------------------------
            # Duplicate check
            # ---------------------------------------------

            key = tuple(
                str(x)
                for x in row[:5]
            )

            if key in existing:
                continue

            ws.append(row)

            existing.add(key)

            added += 1

        print(
            f"Added: {added}"
        )


finally:

    driver.quit()


# =========================================================
# SORT ENTIRE DATASET
# =========================================================

rows = list(
    ws.iter_rows(
        min_row=2,
        values_only=True
    )
)

def safe_date(value):

    if isinstance(value, datetime):
        return value

    text = str(value).strip()

    # Handle Excel/ISO datetime
    try:
        return datetime.fromisoformat(
            text.replace("Z", "")
        )
    except:
        pass

    # Handle DD/MM/YYYY
    try:
        return datetime.strptime(
            text,
            "%d/%m/%Y"
        )
    except:
        pass

    # Last fallback
    try:
        return parser.parse(
            text,
            dayfirst=True
        )
    except:
        return datetime.max


rows.sort(
    key=lambda r: (
        safe_date(r[0]),
        str(r[1]),
        str(r[3])
    )
)

if ws.max_row > 1:
    ws.delete_rows(
        2,
        ws.max_row - 1
    )

for row in rows:
    ws.append(row)


# =========================================================
# FORMAT
# =========================================================

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

for col, width in {
    "A": 14,
    "B": 14,
    "C": 20,
    "D": 10,
    "E": 15,
    "F": 15,
}.items():

    ws.column_dimensions[col].width = width


# =========================================================
# SAVE
# =========================================================

wb.save(FILE)

print()
print("=" * 60)
print("PLANET DATA UPDATE COMPLETE")
print("=" * 60)
print("Previous last year:", last_date.year)
print("Added year:", next_year)
print("File:", FILE)
